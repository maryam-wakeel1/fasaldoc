"""Load local crop records and select compact grounding context for prompts."""

from __future__ import annotations

import json
import re
import shutil
from functools import lru_cache
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import quote

from backend.config import settings

DEFAULT_KNOWLEDGE_PATH = Path(__file__).with_name("crop_knowledge.json")
REFERENCE_IMAGES_DIR = settings.frontend_dir / "assets" / "images"


def _clean_text(value: Any) -> str:
    return str(value or "").strip()


def _as_text_list(value: Any) -> list[str]:
    if isinstance(value, str):
        values = [value]
    elif isinstance(value, (list, tuple)):
        values = value
    else:
        return []
    return [_clean_text(item) for item in values if _clean_text(item)]


def _unique(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        key = value.casefold()
        if value and key not in seen:
            seen.add(key)
            result.append(value)
    return result


def _split_values(value: Any) -> list[str]:
    return _unique([item for item in re.split(r"[\n,;]+", _clean_text(value)) if item.strip()])


def _normalise_reference_path(value: Any) -> str | None:
    raw_path = _clean_text(value).replace("\\", "/")
    if not raw_path or re.match(r"^[A-Za-z]:", raw_path):
        return None

    path = PurePosixPath(raw_path)
    parts = list(path.parts)
    if path.is_absolute() or ".." in parts:
        return None
    if parts and parts[0].casefold() == "images":
        parts = parts[1:]
    if not parts or any(part in {"", ".", ".."} for part in parts):
        return None
    return PurePosixPath(*parts).as_posix()


def _read_sheet_rows(sheet: Any) -> list[dict[str, Any]]:
    header_row = next(sheet.iter_rows(values_only=True), ())
    headers = [_clean_text(value).casefold() for value in header_row]
    rows: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = {
            headers[index]: row[index] if index < len(row) else None
            for index in range(len(headers))
            if headers[index]
        }
        if any(value is not None and _clean_text(value) for value in values.values()):
            rows.append(values)
    return rows


def _mapping_images(workbook: Any) -> dict[tuple[str, str], list[dict[str, str]]]:
    if "Image Mapping" not in workbook.sheetnames:
        return {}

    grouped: dict[tuple[str, str], list[dict[str, str]]] = {}
    for row in _read_sheet_rows(workbook["Image Mapping"]):
        plant_id = _clean_text(row.get("plant_id"))
        disease_id = _clean_text(row.get("disease_id"))
        reference_path = _normalise_reference_path(row.get("image_path"))
        if not plant_id or not disease_id or not reference_path:
            continue
        key = (plant_id, disease_id)
        image = {
            "path": reference_path,
            "image_id": _clean_text(row.get("image_id")),
            "label": _clean_text(row.get("image_label")),
            "status": _clean_text(row.get("status")),
            "review_note": _clean_text(row.get("review_note")),
        }
        images = grouped.setdefault(key, [])
        if reference_path not in {item["path"] for item in images}:
            images.append(image)
    return grouped


def _normalise_enhanced_row(
    row: dict[str, Any], image_mapping: dict[tuple[str, str], list[dict[str, str]]]
) -> dict[str, Any] | None:
    plant_id = _clean_text(row.get("plant_id"))
    disease_id = _clean_text(row.get("disease_id"))
    crop_english = _clean_text(row.get("plant_name"))
    crop_urdu = _clean_text(row.get("plant_name_urdu"))
    disease_english = _clean_text(row.get("disease_name"))
    disease_urdu = _clean_text(row.get("disease_name_urdu"))
    if not crop_english or not disease_english:
        return None

    mapped_images = image_mapping.get((plant_id, disease_id), [])
    dataset_images = [
        reference_path
        for value in _split_values(row.get("image_paths"))
        if (reference_path := _normalise_reference_path(value))
    ]
    references = list(mapped_images)
    known_paths = {image["path"] for image in references}
    for path in dataset_images:
        if path not in known_paths:
            references.append(
                {
                    "path": path,
                    "image_id": "",
                    "label": _clean_text(row.get("image_label")),
                    "status": "",
                    "review_note": "",
                }
            )
            known_paths.add(path)

    disease_aliases = _unique(
        [disease_english, disease_urdu, *_split_values(row.get("aliases"))]
    )
    return {
        "crop": crop_urdu or crop_english,
        "crop_aliases": _unique([crop_english, crop_urdu]),
        "disease": disease_urdu or disease_english,
        "disease_aliases": disease_aliases,
        "symptoms": _clean_text(row.get("symptoms_urdu")) or _clean_text(row.get("symptoms")),
        "low_cost_remedy": _clean_text(row.get("treatment_urdu")) or _clean_text(row.get("treatment")),
        "prevention": _clean_text(row.get("prevention_urdu")) or _clean_text(row.get("prevention")),
        "image_references": [image["path"] for image in references],
        "reference_images": references,
        "plant_id": plant_id,
        "plant_name_en": crop_english,
        "plant_name_urdu": crop_urdu,
        "scientific_name": _clean_text(row.get("scientific_name")),
        "category": _clean_text(row.get("category")),
        "disease_id": disease_id,
        "disease_name_en": disease_english,
        "disease_name_urdu": disease_urdu,
        "disease_type": _clean_text(row.get("disease_type")),
        "causal_agent": _clean_text(row.get("causal_agent")),
        "description_en": _clean_text(row.get("description")),
        "description_urdu": _clean_text(row.get("description_urdu")),
        "symptoms_en": _clean_text(row.get("symptoms")),
        "symptoms_urdu": _clean_text(row.get("symptoms_urdu")),
        "causes_en": _clean_text(row.get("causes")),
        "causes_urdu": _clean_text(row.get("causes_urdu")),
        "treatment_en": _clean_text(row.get("treatment")),
        "treatment_urdu": _clean_text(row.get("treatment_urdu")),
        "prevention_en": _clean_text(row.get("prevention")),
        "prevention_urdu": _clean_text(row.get("prevention_urdu")),
        "source_image_paths": dataset_images,
        "image_label": _clean_text(row.get("image_label")),
        "image_count": row.get("image_count") or len(references),
        "source_name": _clean_text(row.get("source_name")),
        "source_url": _clean_text(row.get("source_url")),
        "confidence": _clean_text(row.get("confidence")),
        "confidence_notes": _clean_text(row.get("confidence_notes")),
    }


def _load_legacy_excel(sheet: Any) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for values in _read_sheet_rows(sheet):
        crop = _clean_text(values.get("crop") or values.get("فصل"))
        disease = _clean_text(values.get("disease") or values.get("بیماری"))
        if not crop or not disease:
            continue
        aliases = _split_values(values.get("crop_aliases") or crop)
        references = [
            reference_path
            for value in _split_values(values.get("image_references") or values.get("images"))
            if (reference_path := _normalise_reference_path(value))
        ]
        rows.append(
            {
                "crop": crop,
                "crop_aliases": aliases,
                "disease": disease,
                "symptoms": _clean_text(values.get("symptoms") or values.get("علامات")),
                "low_cost_remedy": _clean_text(
                    values.get("low_cost_remedy") or values.get("علاج")
                ),
                "prevention": _clean_text(values.get("prevention") or values.get("بچاؤ")),
                "image_references": references,
            }
        )
    return rows


def _load_excel(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Excel import requires openpyxl. Install project requirements first.") from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Dataset"] if "Dataset" in workbook.sheetnames else workbook.active
        records = _read_sheet_rows(sheet)
        if records and {"plant_id", "plant_name", "disease_id", "disease_name"}.issubset(records[0]):
            image_mapping = _mapping_images(workbook)
            return [
                record
                for row in records
                if (record := _normalise_enhanced_row(row, image_mapping)) is not None
            ]
        return _load_legacy_excel(sheet)
    finally:
        workbook.close()


@lru_cache(maxsize=4)
def load_knowledge(source: str | None = None) -> list[dict[str, Any]]:
    path = Path(source) if source else DEFAULT_KNOWLEDGE_PATH
    if path.suffix.lower() == ".xlsx":
        return _load_excel(path)
    with path.open("r", encoding="utf-8") as file:
        data = json.load(file)
    if not isinstance(data, list):
        raise ValueError("Crop knowledge JSON must contain a list of records.")
    return data


def copy_reference_images(images_source: str) -> int:
    source_dir = Path(images_source)
    if not source_dir.is_dir():
        raise ValueError(f"Reference image directory does not exist: {source_dir}")

    source_root = source_dir.resolve()
    REFERENCE_IMAGES_DIR.mkdir(parents=True, exist_ok=True)
    copied = 0
    for source_path in source_dir.rglob("*"):
        if not source_path.is_file():
            continue
        resolved_source = source_path.resolve()
        try:
            relative_path = resolved_source.relative_to(source_root)
        except ValueError:
            continue
        destination = REFERENCE_IMAGES_DIR / relative_path
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(resolved_source, destination)
        copied += 1
    return copied


def import_excel_to_json(
    excel_path: str, output_path: str | None = None, images_source: str | None = None
) -> Path:
    records = _load_excel(Path(excel_path))
    if images_source:
        copy_reference_images(images_source)
    destination = Path(output_path) if output_path else DEFAULT_KNOWLEDGE_PATH
    destination.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    load_knowledge.cache_clear()
    return destination


def _score_records(crop_query: str) -> list[dict[str, Any]]:
    query = (crop_query or "").casefold()
    tokens = [token for token in re.split(r"\s+", query) if len(token) > 1]
    scored: list[tuple[int, dict[str, Any]]] = []
    for record in load_knowledge():
        crop_aliases = _as_text_list(record.get("crop_aliases"))
        disease_aliases = _as_text_list(record.get("disease_aliases"))
        searchable = " ".join(
            [
                _clean_text(record.get("crop")),
                _clean_text(record.get("disease")),
                _clean_text(record.get("symptoms")),
                *crop_aliases,
                *disease_aliases,
            ]
        ).casefold()
        score = sum(1 for token in tokens if token in searchable)
        if any(alias.casefold() in query for alias in crop_aliases if alias):
            score += 3
        if any(alias.casefold() in query for alias in disease_aliases if alias):
            score += 2
        scored.append((score, record))

    selected = [
        record for score, record in sorted(scored, key=lambda item: item[0], reverse=True) if score > 0
    ]
    if not selected and scored:
        selected = [record for _, record in scored[:1]]
    return selected


def get_relevant_records(crop_query: str, max_results: int = 3) -> list[dict[str, Any]]:
    return _score_records(crop_query)[:max_results]


def get_relevant_knowledge(crop_query: str, max_results: int = 3, language: str = "ur") -> str:
    is_english = (language or "ur").strip().lower() == "en"
    records = get_relevant_records(crop_query, max_results)
    if not records:
        return ""

    if is_english:
        blocks = []
        for record in records:
            crop = _clean_text(record.get("plant_name_en")) or _clean_text(record.get("crop"))
            disease = _clean_text(record.get("disease_name_en")) or _clean_text(record.get("disease"))
            symptoms = _clean_text(record.get("symptoms_en")) or _clean_text(record.get("symptoms"))
            remedy = _clean_text(record.get("treatment_en")) or _clean_text(record.get("low_cost_remedy"))
            prevention = _clean_text(record.get("prevention_en")) or _clean_text(record.get("prevention"))
            blocks.append(
                "\n".join(
                    [
                        f"Crop: {crop}",
                        f"Likely disease: {disease}",
                        f"Symptoms: {symptoms}",
                        f"Low-cost guidance: {remedy}",
                        f"Prevention: {prevention}",
                    ]
                )
            )
        return "\n\n".join(blocks)

    return "\n\n".join(
        "\n".join(
            [
                f"فصل: {record['crop']}",
                f"ممکنہ بیماری: {record['disease']}",
                f"علامات: {record['symptoms']}",
                f"کم خرچ رہنمائی: {record['low_cost_remedy']}",
                f"بچاؤ: {record['prevention']}",
            ]
        )
        for record in records
    )


def resolve_reference_images(record: dict[str, Any]) -> list[dict[str, Any]]:
    try:
        root = REFERENCE_IMAGES_DIR.resolve()
    except OSError:
        return []

    references = record.get("reference_images") or [
        {"path": path} for path in _as_text_list(record.get("image_references"))
    ]
    resolved_images: list[dict[str, Any]] = []
    seen: set[str] = set()
    needs_verification = _clean_text(record.get("confidence")).casefold() == "needs_verification"
    caption_urdu = f"{_clean_text(record.get('crop'))} — {_clean_text(record.get('disease'))}".strip(" —")
    english_crop = _clean_text(record.get("plant_name_en")) or _clean_text(record.get("crop"))
    english_disease = _clean_text(record.get("disease_name_en")) or _clean_text(record.get("disease"))
    caption_english = f"{english_crop} — {english_disease}".strip(" —")
    for reference in references:
        raw_path = reference.get("path") if isinstance(reference, dict) else reference
        relative_path = _normalise_reference_path(raw_path)
        if not relative_path or relative_path in seen:
            continue
        candidate = REFERENCE_IMAGES_DIR / relative_path
        try:
            resolved_path = candidate.resolve()
            resolved_path.relative_to(root)
        except (OSError, ValueError):
            continue
        if not resolved_path.is_file():
            continue
        seen.add(relative_path)
        resolved_images.append(
            {
                "url": f"/static/assets/images/{quote(relative_path, safe='/')}",
                "crop": _clean_text(record.get("crop")),
                "disease": _clean_text(record.get("disease")),
                "caption_urdu": caption_urdu,
                "caption_english": caption_english,
                "needs_verification": needs_verification,
            }
        )
    return resolved_images


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Convert crop disease Excel data into FasalDoc JSON.")
    parser.add_argument("excel_path")
    parser.add_argument("--output")
    parser.add_argument("--images-source")
    arguments = parser.parse_args()
    print(import_excel_to_json(arguments.excel_path, arguments.output, arguments.images_source))
